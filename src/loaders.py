
from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document
from openpyxl import load_workbook
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".xlsx",
}


def load_document(file_path: str | Path) -> list[Document]:
    """
    Load a supported file and convert it into LangChain Documents.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(
            f"Unsupported file type: {extension}. "
            f"Supported types: {supported}"
        )

    if extension == ".pdf":
        documents = _load_pdf(path)

    elif extension == ".docx":
        documents = _load_docx(path)

    elif extension in {".txt", ".md"}:
        documents = _load_text(path)

    elif extension == ".xlsx":
        documents = _load_xlsx(path)

    else:
        raise ValueError(f"Unsupported file type: {extension}")

    for document in documents:
        document.metadata["source"] = path.name
        document.metadata["file_type"] = extension.lstrip(".")

    return documents


def _load_pdf(path: Path) -> list[Document]:
    """Load a PDF one page at a time."""

    reader = PdfReader(str(path))
    documents = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        documents.append(
            Document(
                page_content=text,
                metadata={
                    "page": page_number,
                },
            )
        )

    return documents


def _load_docx(path: Path) -> list[Document]:
    """Load paragraphs from a DOCX file."""

    document = DocxDocument(str(path))

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    text = "\n".join(paragraphs)

    return [
        Document(
            page_content=text,
            metadata={},
        )
    ]


def _load_text(path: Path) -> list[Document]:
    """Load TXT or Markdown as a single document."""

    text = path.read_text(encoding="utf-8")

    return [
        Document(
            page_content=text,
            metadata={},
        )
    ]


def _load_xlsx(path: Path) -> list[Document]:
    """Load an XLSX workbook, preserving sheet boundaries."""

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    documents = []

    for worksheet in workbook.worksheets:
        rows = []

        for row in worksheet.iter_rows(values_only=True):
            values = [
                str(value).strip()
                for value in row
                if value is not None
            ]

            if values:
                rows.append(" | ".join(values))

        if rows:
            documents.append(
                Document(
                    page_content="\n".join(rows),
                    metadata={
                        "sheet": worksheet.title,
                    },
                )
            )

    workbook.close()

    return documents